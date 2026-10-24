"""Kiểm thử luồng đọc file -> chấm -> báo cáo (không cần API key, không cần mạng)."""
import io
import os
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
os.environ["CV_AGENT_AS_OF"] = "2026-10"

from cv_agent import config, io_utils
from cv_agent.facts import extract_facts
from cv_agent.pipeline import build_rubric, get_scorer, list_cv_files, screen
from cv_agent.report import rank, suitable, write_csv, write_excel
from cv_agent.verdict import ERROR

JD = """Yêu cầu:
- Lập trình Python tốt.
- Biết PyTorch hoặc TensorFlow.
- Biết sử dụng Git.
"""
GOOD = """Nguyễn Văn A
AI Engineer
Email: a@example.com • 0903 456 789 • github.com/someone

EXPERIENCE
AI Engineer – Acme (03/2021 - Present)
- Built models with Python and PyTorch, versioned with Git on GitHub.
"""
WEAK = """Trần Thị B
Accountant

EXPERIENCE
Accountant – Foo (2020 - 2024)
- Prepared tax reports in Excel.
"""


class Workdir(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def write(self, name, text, encoding="utf-8"):
        p = self.dir / name
        p.write_text(text, encoding=encoding)
        return p


class ReadingTests(Workdir):
    def test_encoding_detection(self):
        self.assertEqual(io_utils.decode_bytes("Kinh nghiệm".encode("utf-8")), "Kinh nghiệm")
        self.assertEqual(io_utils.decode_bytes("Kinh nghiệm".encode("utf-16")), "Kinh nghiệm")

    def test_ansi_file_is_not_misread_as_utf16(self):
        raw = "Nguyen Van A - Python".encode("ascii") + b"\xe9"           # độ dài chẵn, không BOM: trước đây có thể bị giải mã thành chữ Hán
        self.assertNotRegex(io_utils.decode_bytes(raw), r"[一-鿿]")

    def test_clear_errors(self):
        with self.assertRaisesRegex(ValueError, "Chưa hỗ trợ"):
            io_utils.read_document(self.write("x.doc", "abc"))
        with self.assertRaisesRegex(ValueError, "không tìm thấy"):
            io_utils.read_document(self.dir / "missing.txt")
        bad = self.dir / "bad.pdf"
        bad.write_bytes(b"%PDF-1.4 broken")
        with self.assertRaisesRegex(ValueError, "không mở được PDF"):
            io_utils.read_document(bad)

    def test_huge_text_is_truncated_with_warning(self):
        p = self.write("big.txt", "python " * 100_000)
        doc = io_utils.read_document(p)
        self.assertEqual(len(doc.text), config.MAX_TEXT_CHARS)
        self.assertTrue(doc.warnings)

    def test_two_column_pdf_is_read_column_by_column(self):
        import pymupdf
        pdf = pymupdf.open()
        page = pdf.new_page()
        page.insert_text((150, 50), "NGUYEN VAN A", fontsize=16)
        for i, t in enumerate(["SKILLS", "Python", "PyTorch", "Docker", "Git"] * 6):
            page.insert_text((40, 100 + i * 18), t, fontsize=10)
        for i, t in enumerate(["EXPERIENCE", "ML Engineer - Acme (03/2021 - Present)", "Built models", "EDUCATION", "B.Sc. (2016 - 2020)"] * 6):
            page.insert_text((260, 100 + i * 18), t, fontsize=10)
        path = self.dir / "two.pdf"
        pdf.save(str(path))
        lines = io_utils.read_text(path).splitlines()
        self.assertEqual(lines[0], "NGUYEN VAN A")                                   # tiêu đề trải ngang được giữ nguyên vẹn
        last_left = max(i for i, l in enumerate(lines) if l in ("Python", "PyTorch", "Docker", "Git"))
        first_right = min(i for i, l in enumerate(lines) if l.startswith("EXPERIENCE"))
        self.assertLess(last_left, first_right)                                      # hết cột trái rồi mới sang cột phải

    def test_pdf_bullet_glyph_artifacts_are_cleaned(self):
        cleaned = io_utils._clean_pdf_text("EXPERIENCE\n(cid:127) Built models\nText (cid:128) mid\n")
        self.assertNotIn("cid", cleaned)
        self.assertIn("- Built models", cleaned)                                       # ký hiệu đầu dòng thành gạch đầu dòng

    def test_ocr_spacing_fix_keeps_tech_names(self):
        fixed = io_utils.fix_ocr_spacing("ProductRecommendationService2023-2024 inPyTorch,rate ServedviaFastAPIandDockeronAWSECS")
        self.assertIn("Product Recommendation Service 2023-2024", fixed)
        self.assertIn("PyTorch", fixed)
        self.assertIn("FastAPI", fixed)


class FactsTests(unittest.TestCase):
    def test_contact_and_gpa(self):
        f = extract_facts(GOOD + "\nGPA 3.2 / 4.0\nTOEIC 790")
        self.assertEqual(f["candidate_name"], "Nguyễn Văn A")
        self.assertEqual(f["email"], "a@example.com")
        self.assertEqual(f["gpa"], "3.2/4")
        self.assertEqual(f["english"], "TOEIC 790")
        self.assertIn("github.com/someone", f["github"])

    def test_phone_gpa_and_title_formats(self):
        for text, expected in (("Phone: +84 (0)912 345 678", "+84 (0)912 345 678"), ("ĐT: 0912.345.678", "0912.345.678"),
                               ("Tel: 024 3825 1234", "024 3825 1234"), ("SĐT: (+84) 912-345-678", "(+84) 912-345-678")):
            self.assertEqual(extract_facts(text)["phone"], expected, text)
        self.assertEqual(extract_facts("Ngày 20210123456 mã số")["phone"], "")                       # số dài không phải điện thoại
        for text, expected in (("GPA: 8.5/10", "8.5/10"), ("DTB 2.9/4", "2.9/4"), ("Điểm trung bình: 7.8", "7.8/10"), ("GPA: 85/100", "85/100"),
                               ("CGPA: 3,2/4", "3.2/4")):
            self.assertEqual(extract_facts(text)["gpa"], expected, text)
        self.assertEqual(extract_facts("Mr. John Smith\nEngineer\n")["candidate_name"], "John Smith")
        self.assertEqual(extract_facts("Ths. Nguyễn Văn An\nGiám đốc kỹ thuật\n")["candidate_name"], "Nguyễn Văn An")

    def test_missing_facts_are_empty_not_crash(self):
        f = extract_facts("chỉ có vài chữ")
        self.assertEqual((f["email"], f["gpa"], f["github"]), ("", "", ""))


class PipelineTests(Workdir):
    def setUp(self):
        super().setUp()
        from unittest import mock
        import cv_agent.pipeline as pl
        patcher = mock.patch.object(pl, "MIN_CV_WORDS", 0)          # CV mẫu trong test rất ngắn; ngưỡng thật kiểm riêng ở test_too_little_text
        patcher.start()
        self.addCleanup(patcher.stop)
        self.rubric = build_rubric(JD)

    def items(self, **files):
        return [(n, self.write(n, t)) for n, t in files.items()]

    def test_ranking_and_suitable_list(self):
        rows = screen(self.rubric, self.items(**{"good.txt": GOOD, "weak.txt": WEAK}), workers=2)
        self.assertEqual([r["file"] for r in rows], ["good.txt", "weak.txt"])
        self.assertEqual([r["rank"] for r in rows], [1, 2])
        self.assertEqual(rows[0]["candidate_name"], "Nguyễn Văn A")
        self.assertIn("good.txt", [r["file"] for r in suitable(rows)])
        self.assertNotIn("weak.txt", [r["file"] for r in suitable(rows)])

    def test_one_bad_file_does_not_break_the_batch(self):
        items = self.items(**{"good.txt": GOOD, "empty.txt": ""}) + [("ghost.pdf", self.dir / "ghost.pdf")]
        rows = screen(self.rubric, items)
        verdicts = {r["file"]: r["verdict"] for r in rows}
        self.assertEqual(verdicts["empty.txt"], ERROR)
        self.assertEqual(verdicts["ghost.pdf"], ERROR)
        self.assertNotEqual(verdicts["good.txt"], ERROR)
        self.assertEqual(rows[-1]["verdict"], ERROR)                                 # CV lỗi luôn ở cuối
        for r in rows:
            if r["verdict"] == ERROR:
                self.assertEqual(r["needs_human_review"], "CÓ")

    def test_duplicate_cv_is_flagged(self):
        rows = screen(self.rubric, self.items(**{"a.txt": GOOD, "b.txt": GOOD}), workers=1)
        dup = [r for r in rows if "trùng" in r["review_reason"]]
        self.assertEqual(len(dup), 1)
        self.assertEqual(dup[0]["file"], "b.txt")                                    # bản đứng sau bị đánh dấu, bản đầu giữ nguyên
        self.assertEqual(dup[0]["needs_human_review"], "CÓ")

    def test_bad_thresholds_rejected_early(self):
        with self.assertRaises(ValueError):
            screen(self.rubric, self.items(**{"a.txt": GOOD}), fit=40, maybe=50, reject=30)

    def test_invalid_dates_are_flagged_not_silent(self):
        cv = "A B\nEXPERIENCE\nEngineer – X (2023 - 2019)\n- Python PyTorch Git\n"
        row = screen(self.rubric, self.items(**{"d.txt": cv}))[0]
        self.assertIn("không hợp lý", row["red_flags"])

    def test_scanned_pdf_is_flagged_as_ocr(self):
        import pymupdf
        from unittest import mock
        src = pymupdf.open()
        src.new_page().insert_text((72, 100), "x")
        png = src[0].get_pixmap(dpi=50).tobytes("png")
        out = pymupdf.open()
        out.new_page().insert_image(pymupdf.Rect(0, 0, 595, 842), stream=png)
        path = self.dir / "scan.pdf"
        out.save(str(path))
        with mock.patch.object(io_utils, "_ocr_tesseract", return_value=GOOD):
            row = screen(self.rubric, [("scan.pdf", path)])[0]
        self.assertEqual(row["source"], "ocr:tesseract")
        self.assertIn("OCR", row["review_reason"])
        self.assertEqual(row["needs_human_review"], "CÓ")

    def test_too_little_text_is_an_error_not_a_silent_reject(self):
        import cv_agent.pipeline as pl
        from unittest import mock
        with mock.patch.object(pl, "MIN_CV_WORDS", 30):
            row = screen(self.rubric, self.items(**{"tiny.txt": "Nguyễn Văn A\nPython và Git trong 3 năm làm việc ở công ty."}))[0]
        self.assertEqual(row["verdict"], ERROR)
        self.assertIn("không đủ để đánh giá", row["review_reason"])

    def test_list_cv_files_skips_temp_and_hidden(self):
        for n in ("a.txt", "~$b.docx", ".hidden.txt", "c.xlsx", "d.pdf"):
            self.write(n, "x")
        self.assertEqual([f.name for f in list_cv_files(self.dir)], ["a.txt", "d.pdf"])


class ReportTests(Workdir):
    def setUp(self):
        super().setUp()
        from unittest import mock
        import cv_agent.pipeline as pl
        patcher = mock.patch.object(pl, "MIN_CV_WORDS", 0)
        patcher.start()
        self.addCleanup(patcher.stop)

    def test_excel_and_csv_are_readable(self):
        rubric = build_rubric(JD)
        items = [("good.txt", self.write("good.txt", GOOD)), ("weak.txt", self.write("weak.txt", WEAK))]
        rows = screen(rubric, items)
        import openpyxl
        buf = io.BytesIO()
        write_excel(rows, buf)
        wb = openpyxl.load_workbook(io.BytesIO(buf.getvalue()))
        self.assertEqual(wb.sheetnames, ["Danh sách phù hợp", "Tất cả CV"])
        self.assertEqual(wb["Tất cả CV"].max_row, 3)
        header = [c.value for c in wb["Tất cả CV"][1]]
        for needed in ("Ứng viên", "Email", "SĐT", "Kết luận", "Vì sao kết luận vậy", "HR cần xem gì"):
            self.assertIn(needed, header)
        csv_path = self.dir / "r.csv"
        write_csv(rows, csv_path)
        self.assertIn("Nguyễn Văn A", csv_path.read_text(encoding="utf-8-sig"))

    def test_rank_breaks_ties_by_core_ratio(self):
        rows = rank([{"overall_score": 50, "core_ratio": 0.5, "verdict": "x", "file": "a"},
                     {"overall_score": 50, "core_ratio": 0.9, "verdict": "x", "file": "b"},
                     {"overall_score": -1, "core_ratio": 0.0, "verdict": ERROR, "file": "c"}])
        self.assertEqual([r["file"] for r in rows], ["b", "a", "c"])


class LlmIsOptionalTests(unittest.TestCase):
    def test_llm_engine_is_off_by_default(self):
        if config.LLM_ENABLED:
            self.skipTest("LLM đã được bật trong môi trường này")
        with self.assertRaises(RuntimeError):
            get_scorer("llm")
        with self.assertRaises(RuntimeError):
            build_rubric(JD, "llm")

    def test_core_modules_do_not_import_llm_sdk(self):
        import subprocess
        code = ("import sys; import cv_agent.pipeline, cv_agent.report, cv_agent.baseline, cv_agent.verdict; "
                "bad=[m for m in sys.modules if m.split('.')[0] in ('anthropic','openai')]; print(bad)")
        out = subprocess.run([sys.executable, "-c", code], capture_output=True, text=True,
                             cwd=str(Path(__file__).resolve().parent.parent))
        self.assertEqual(out.stdout.strip(), "[]", out.stderr)


class SecurityAndFormatTests(Workdir):
    def test_formula_in_cv_text_is_not_executed_in_excel_or_csv(self):
        import openpyxl
        from cv_agent.report import COLUMNS
        row = {k: "" for k, _, _ in COLUMNS}
        row.update(verdict="Loại", needs_human_review="Không", overall_score=1, candidate_name='=HYPERLINK("http://evil","x")',
                   email="-1+1", file="a.txt", rank=1)
        buf = io.BytesIO()
        write_excel([row], buf)
        ws = openpyxl.load_workbook(io.BytesIO(buf.getvalue()))["Tất cả CV"]
        formulas = [c for c in ws[2] if c.data_type == "f"]
        self.assertEqual(formulas, [])
        path = self.dir / "r.csv"
        write_csv([row], path)
        line = path.read_text(encoding="utf-8-sig").splitlines()[1]
        self.assertIn("'=HYPERLINK", line)
        self.assertIn("'-1+1", line)

    def test_control_characters_and_huge_cells_do_not_break_excel(self):
        import openpyxl
        from cv_agent.report import COLUMNS
        row = {k: "" for k, _, _ in COLUMNS}
        row.update(verdict="Loại", needs_human_review="Không", overall_score=1, candidate_name="Nguyễn\x07 Văn\x00 A",
                   gaps="x" * 50000, file="a.txt", rank=1)
        buf = io.BytesIO()
        write_excel([row], buf)                                                         # trước đây: IllegalCharacterError
        ws = openpyxl.load_workbook(io.BytesIO(buf.getvalue()))["Tất cả CV"]
        header = [c.value for c in ws[1]]
        self.assertEqual(ws.cell(2, header.index("Ứng viên") + 1).value, "Nguyễn Văn A")
        self.assertLessEqual(len(ws.cell(2, header.index("Khoảng trống") + 1).value), 32001)

    def test_docx_header_table_and_columns_are_read(self):
        import docx
        d = docx.Document()
        d.sections[0].header.paragraphs[0].text = "Nguyễn Văn A | a@example.com"
        d.add_paragraph("EXPERIENCE")
        t = d.add_table(rows=1, cols=2)
        t.rows[0].cells[0].text = "SKILLS\nPython, PyTorch"
        t.rows[0].cells[1].text = "EDUCATION\nB.Sc. (2016 - 2020)"
        path = self.dir / "cv.docx"
        d.save(str(path))
        lines = io_utils.read_text(path).splitlines()
        self.assertEqual(lines[0], "Nguyễn Văn A | a@example.com")                     # đầu trang được đọc
        self.assertIn("SKILLS", lines)                                                  # xuống dòng mềm trong ô bảng thành dòng riêng
        self.assertLess(lines.index("Python, PyTorch"), lines.index("EDUCATION"))      # hết cột trái rồi mới sang cột phải

    def test_missing_end_of_file_does_not_break(self):
        for name in ("a.txt", "b.md"):
            self.assertEqual(io_utils.read_text(self.write(name, "x" * 10)), "x" * 10)


if __name__ == "__main__":
    unittest.main()
