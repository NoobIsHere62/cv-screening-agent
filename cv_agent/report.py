"""Gom kết quả chấm 1 CV thành 1 dòng báo cáo (dùng chung cho run.py và app.py) và xuất CSV/Excel."""
from .config import CORE_MIN, FIT_THRESHOLD, MAYBE_THRESHOLD, REJECT_THRESHOLD
from .verdict import COLORS, ERROR, FIT, MAYBE, REJECT, RESERVE, SUITABLE, verdict

COLUMNS = [  # (khoá, tiêu đề, độ rộng)
    ("rank", "Hạng", 6), ("candidate_name", "Ứng viên", 22), ("email", "Email", 26), ("phone", "SĐT", 15),
    ("overall_score", "Điểm", 8), ("required_score", "% đáp ứng yêu cầu bắt buộc", 12), ("verdict", "Kết luận", 24),
    ("suggestion", "Đề xuất", 14), ("checklist_summary", "Đối chiếu yêu cầu", 38), ("verdict_reason", "Vì sao kết luận vậy", 44),
    ("needs_human_review", "Cần người xem lại?", 12), ("review_reason", "HR cần xem gì", 60),
    ("hr_decision", "Quyết định của HR", 16), ("hr_note", "Ghi chú của HR", 30),
    ("total_years", "Kinh nghiệm (năm, thực tập tính ½)", 14), ("gpa", "GPA", 8), ("english", "Ngoại ngữ", 12),
    ("github", "GitHub", 26), ("strengths", "Điểm mạnh", 45), ("gaps", "Khoảng trống", 45),
    ("red_flags", "Cảnh báo", 40), ("failed_hard", "Trượt điều kiện cứng", 36), ("unclear_hard", "Chưa rõ điều kiện cứng", 36),
    ("decision", "Kết luận 3 mức", 14), ("file", "File", 28), ("note", "Ghi chú lỗi", 40),
]


_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def csv_safe(value):
    """Chống "CSV/Excel injection": ô văn bản từ CV bắt đầu bằng = + - @ có thể bị Excel chạy như công thức. Thêm dấu nháy đơn."""
    if isinstance(value, str) and value.startswith(_FORMULA_START):
        return "'" + value
    return value


def _hard_lines(rubric, ids, notes) -> str:
    desc = {h.id: h.description for h in rubric.hard_requirements}
    return " | ".join(f"{desc.get(i, i)}" + (f": {notes[i]}" if notes.get(i) else "") for i in ids)


def review_reasons(rubric, s, d, st, doc=None) -> list[str]:
    """Các việc cụ thể HR cần kiểm tra trên CV này (rỗng = không cần xem lại)."""
    hard = {h.id: h for h in rubric.hard_requirements}
    notes = {h.requirement_id: h.note for h in d.hard_checks}
    names = {c.id: c.name for c in rubric.criteria}
    out = []
    if doc is not None:
        out += list(doc.warnings)
    for hid in s["unclear_hard"]:
        h = hard.get(hid)
        out.append(f"Chưa xác minh được điều kiện «{h.description if h else hid}»: {notes.get(hid) or 'CV không có thông tin'}")
    for cid in s["missing_criteria"]:
        out.append(f"Chưa chấm được tiêu chí «{names.get(cid, cid)}», cần đọc CV để đánh giá")
    if st["evidence_dropped"]:
        out.append(f"{st['evidence_dropped']} trích dẫn bằng chứng không khớp CV, nên kiểm tra lại điểm")
    out += [f"Cảnh báo: {rf}" for rf in d.red_flags]
    return out


SUGGESTION = {FIT: "Mời phỏng vấn", MAYBE: "Cân nhắc mời", RESERVE: "Giữ dự phòng", REJECT: "Không mời", ERROR: "Xem thủ công"}
HR_FROM_SUGGESTION = {FIT: "Mời phỏng vấn", MAYBE: "Giữ lại", RESERVE: "Giữ lại", REJECT: "Từ chối"}   # khi HR bấm "chấp nhận đề xuất"


def _status_of_score(score) -> str:
    return "unclear" if score is None else "strong" if score >= 3 else "partial" if score == 2 else "weak" if score == 1 else "missing"


def build_checklist(rubric, d) -> list[dict]:
    """Bảng đối chiếu từng yêu cầu của JD với CV: điều kiện cứng (met/failed/unclear) và từng kỹ năng (strong/partial/weak/missing).
    HR quyết định theo danh sách này dễ hơn là theo một con số."""
    hard = {h.requirement_id: h for h in d.hard_checks}
    scores = {c.criterion_id: c for c in d.criterion_scores}
    items = []
    for h in rubric.hard_requirements:
        hc = hard.get(h.id)
        status = {"met": "met", "not_met": "failed", "unclear": "unclear"}.get(hc.status if hc else "unclear", "unclear")
        items.append({"kind": "hard", "type": h.type, "label": h.source_quote, "status": status,
                      "note": hc.note if hc else "", "evidence": [hc.evidence] if hc and hc.evidence else []})
    year_quotes = {h.source_quote for h in rubric.hard_requirements if h.type == "experience_years"}
    for c in rubric.criteria:
        if c.source_quote in year_quotes:                  # tiêu chí số năm đã hiện ở phần điều kiện bắt buộc, không lặp lại ở kỹ năng
            continue
        cs = scores.get(c.id)
        score = cs.score if cs else None
        items.append({"kind": "required" if c.importance == "required" else "preferred", "label": c.source_quote or c.name,
                      "score": score, "status": _status_of_score(score), "note": cs.reasoning if cs else "",
                      "evidence": list(cs.evidence[:3]) if cs else []})
    return items


def checklist_counts(items: list[dict]) -> dict:
    def count(kind, ok):
        sel = [i for i in items if i["kind"] == kind]
        return sum(i["status"] in ok for i in sel), len(sel)
    hard_met, hard_total = count("hard", {"met"})
    req_met, req_total = count("required", {"strong", "partial"})
    pref_met, pref_total = count("preferred", {"strong", "partial"})
    return {"hard_met": hard_met, "hard_total": hard_total, "req_met": req_met, "req_total": req_total,
            "pref_met": pref_met, "pref_total": pref_total}


def checklist_summary(c: dict) -> str:
    parts = []
    if c["hard_total"]:
        parts.append(f"{c['hard_met']}/{c['hard_total']} điều kiện bắt buộc")
    if c["req_total"]:
        parts.append(f"{c['req_met']}/{c['req_total']} kỹ năng bắt buộc có bằng chứng")
    if c["pref_total"]:
        parts.append(f"{c['pref_met']}/{c['pref_total']} điểm cộng")
    return " · ".join(parts)


def verdict_reason(rubric, s, notes: dict, v: str, fit, maybe, reject) -> str:
    """Một câu giải thích vì sao CV nhận kết luận này (HR hay hỏi nhất về nhóm 'Loại'). notes: {mã điều kiện cứng: ghi chú}."""
    score = s.get("required_score", s["overall_score"])
    if v == REJECT:
        if s["failed_hard"]:
            return "Trượt điều kiện bắt buộc: " + _hard_lines(rubric, s["failed_hard"], notes)
        return f"Chỉ đáp ứng {score:g}% yêu cầu bắt buộc, thấp hơn mức tối thiểu {reject:g}%"
    if v == RESERVE and s.get("core_ratio", 1.0) < CORE_MIN:
        return f"Mới có bằng chứng cho {s['core_ratio']:.0%} kỹ năng bắt buộc (cần ≥{CORE_MIN:.0%}) nên chỉ ở mức dự phòng"
    if v == FIT:
        return f"Đáp ứng {score:g}% yêu cầu bắt buộc (cần ≥{fit:g}%) và đạt đủ điều kiện cứng"
    if v == MAYBE:
        if score >= fit:
            open_items = []
            if s["unclear_hard"]:
                hard = {h.id: h.description for h in rubric.hard_requirements}
                open_items.append("điều kiện " + ", ".join(f"«{hard.get(i, i)}»" for i in s["unclear_hard"]))
            if s["required_zero"]:
                open_items.append(f"{len(s['required_zero'])} kỹ năng bắt buộc chưa thấy bằng chứng")
            return f"Đáp ứng {score:g}% yêu cầu bắt buộc (đủ mức phù hợp) nhưng còn " + (" và ".join(open_items) or "điểm cần HR xác nhận")
        return f"Đáp ứng {score:g}% yêu cầu bắt buộc (mức {maybe:g}–{fit:g}%): có thể mời nếu HR xác nhận được phần còn thiếu"
    return f"Đáp ứng {score:g}% yêu cầu bắt buộc (mức {reject:g}–{maybe:g}%): chưa đủ để mời, giữ lại dự phòng"


def join_reasons(items: list[str]) -> str:
    return "\n".join(f"{i}. {r}" for i, r in enumerate(items, 1))


def make_row(rubric, name: str, out: dict, fit=FIT_THRESHOLD, maybe=MAYBE_THRESHOLD, reject=REJECT_THRESHOLD,
             doc=None, facts=None) -> dict:
    s, d, st = out["summary"], out["detail"], out["stats"]
    names = {c.id: c.name for c in rubric.criteria}
    notes = {h.requirement_id: h.note for h in d.hard_checks}
    facts = facts or {}
    reasons = review_reasons(rubric, s, d, st, doc)
    src = facts.get("name_source", "")
    if facts and not facts.get("candidate_name"):
        reasons.append("Không tìm thấy tên ứng viên trong CV, đang hiển thị tên file")
    elif src in ("file", "email", "ocr"):
        where = {"file": "tên file", "email": "địa chỉ email", "ocr": "chữ OCR (có thể sai chính tả hoặc mất dấu)"}[src]
        reasons.append(f"Tên ứng viên lấy từ {where}, cần xác nhận")
    v = verdict(s, fit, maybe, reject)
    years = facts.get("total_years")
    checklist = build_checklist(rubric, d)
    counts = checklist_counts(checklist)
    return {
        "file": name,
        "candidate_name": facts.get("candidate_name") or name.rsplit(".", 1)[0],
        "name_source": src,
        "email": facts.get("email", ""), "phone": facts.get("phone", ""), "gpa": facts.get("gpa", ""),
        "english": facts.get("english", ""), "github": facts.get("github", ""),
        "total_years": "" if years is None else years,
        "overall_score": s["overall_score"],
        "required_score": s.get("required_score", s["overall_score"]),
        "core_ratio": s.get("core_ratio", 1.0),
        "verdict": v,
        "suggestion": SUGGESTION[v],
        "verdict_reason": verdict_reason(rubric, s, notes, v, fit, maybe, reject),
        # đầu vào của kết luận: đủ để đổi ngưỡng sau này mà không phải đọc/chấm lại CV (xem rejudge)
        "verdict_inputs": {k: s.get(k) for k in _VERDICT_KEYS}, "hard_notes": {hid: note for hid, note in notes.items() if note},
        "checklist": checklist,
        "checklist_summary": checklist_summary(counts),
        **counts,
        "decision": s["decision"],
        "failed_hard": _hard_lines(rubric, s["failed_hard"], notes),
        "unclear_hard": _hard_lines(rubric, s["unclear_hard"], notes),
        "needs_human_review": "CÓ" if reasons else "Không",
        "review_items": reasons,
        "review_reason": join_reasons(reasons),
        "strengths": " | ".join(d.strengths),
        "gaps": " | ".join(d.gaps),
        "red_flags": " | ".join(d.red_flags),
        "evidence_dropped": f'{st["evidence_dropped"]}/{st["evidence_total"]}',
        "source": doc.source if doc is not None else "text",
        # phần chi tiết cho giao diện (không ghi vào CSV/Excel)
        "criteria": [(names.get(c.criterion_id, c.criterion_id), c.score, c.reasoning, c.evidence)
                     for c in d.criterion_scores],
    }


_VERDICT_KEYS = ("overall_score", "required_score", "core_ratio", "failed_hard", "unclear_hard", "required_zero", "missing_criteria")


def can_rejudge(rows: list[dict]) -> bool:
    """Các lần chạy lưu bằng phiên bản cũ không có `verdict_inputs`, nên không đổi ngưỡng được."""
    return all("verdict_inputs" in r for r in rows if r.get("verdict") != ERROR)


def rejudge(rubric, rows: list[dict], fit: float, maybe: float, reject: float) -> list[dict]:
    """Tính lại kết luận / đề xuất / lý do theo ngưỡng mới và xếp hạng lại, giữ nguyên điểm, bằng chứng và quyết định của HR.
    Sửa trực tiếp `rows` và trả về chính nó."""
    for r in rows:
        if r.get("verdict") == ERROR:
            continue
        s = r["verdict_inputs"]
        v = verdict(s, fit, maybe, reject)
        r["verdict"], r["suggestion"] = v, SUGGESTION[v]
        r["verdict_reason"] = verdict_reason(rubric, s, r.get("hard_notes", {}), v, fit, maybe, reject)
    return rank(rows)


def error_row(name: str, err) -> dict:
    return {"file": name, "candidate_name": name.rsplit(".", 1)[0], "verdict": ERROR, "decision": ERROR, "suggestion": SUGGESTION[ERROR],
            "note": str(err), "overall_score": -1, "required_score": -1, "core_ratio": 0.0, "needs_human_review": "CÓ",
            "verdict_reason": "Không đọc/chấm được CV", "checklist": [], "checklist_summary": "",
            "review_items": [f"Không đọc/chấm được CV này, cần xem thủ công: {err}"],
            "review_reason": f"1. Không đọc/chấm được CV này, cần xem thủ công: {err}"}

def rank(rows: list[dict]) -> list[dict]:
    """Nhóm kết luận trước (Phù hợp > Có thể mời > Dự phòng > Loại > Lỗi); trong nhóm theo điểm tổng (có điểm cộng của kỹ năng
    ưu tiên), bằng nhau thì theo mức đáp ứng yêu cầu bắt buộc rồi tỉ lệ kỹ năng bắt buộc."""
    tier = {FIT: 0, MAYBE: 1, RESERVE: 2, REJECT: 3, ERROR: 4}
    rows.sort(key=lambda r: (-tier.get(r.get("verdict"), 4), r.get("overall_score", -1),
                             r.get("required_score", -1), r.get("core_ratio", 0.0)), reverse=True)
    for i, r in enumerate(rows, 1):
        r["rank"] = i
    return rows


def suitable(rows: list[dict]) -> list[dict]:
    """Danh sách phù hợp (đã xếp hạng): Phù hợp - Mời phỏng vấn + Có thể mời phỏng vấn."""
    return [r for r in rows if r["verdict"] in SUITABLE]


def count_by_verdict(rows: list[dict]) -> dict:
    return {v: sum(r["verdict"] == v for r in rows) for v in (FIT, MAYBE, RESERVE, REJECT, ERROR)}


def write_excel(rows: list[dict], out):
    """`out` là đường dẫn hoặc file-like (BytesIO). Sheet 1: danh sách phù hợp; sheet 2: tất cả CV."""
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    def cell_value(v):
        """Ký tự điều khiển làm openpyxl báo lỗi, ô quá 32.767 ký tự làm Excel báo file hỏng: làm sạch và cắt bớt."""
        if isinstance(v, str):
            v = ILLEGAL_CHARACTERS_RE.sub("", v)
            return v if len(v) <= 32000 else v[:32000] + "…"
        return v

    wb = Workbook()
    keys = [k for k, _, _ in COLUMNS]
    col_verdict, col_flag, col_reason = (keys.index(k) + 1 for k in ("verdict", "needs_human_review", "review_reason"))

    wrap = Alignment(wrap_text=True, vertical="top")                    # dùng chung một đối tượng kiểu cho mọi ô (tạo mới cho từng ô rất chậm)
    verdict_fill = {v: PatternFill("solid", fgColor=c) for v, c in COLORS.items()}
    review_fill, review_font = PatternFill("solid", fgColor="F8CBAD"), Font(bold=True, color="C00000")

    def fill_sheet(ws, data):
        ws.append([title for _, title, _ in COLUMNS])
        for i, (_, _, width) in enumerate(COLUMNS, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
            ws.cell(1, i).font = Font(bold=True)
        for row, r in enumerate(data, 2):                               # đếm hàng tại chỗ: ws.max_row duyệt mọi ô mỗi lần gọi, làm cả vòng lặp O(n²)
            ws.append([cell_value("" if r.get(k, "") == -1 else r.get(k, "")) for k in keys])
            for cell in ws[row]:
                cell.alignment = wrap
                if isinstance(cell.value, str) and cell.value.startswith("="):      # văn bản từ CV không bao giờ được là công thức
                    cell.data_type = "s"
            if r["verdict"] in verdict_fill:
                ws.cell(row, col_verdict).fill = verdict_fill[r["verdict"]]
            if r.get("needs_human_review") == "CÓ":                     # nổi bật ô cần xem
                for col in (col_flag, col_reason):
                    ws.cell(row, col).fill = review_fill
                ws.cell(row, col_flag).font = review_font
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(len(data) + 1, 1)}"

    fill_sheet(wb.active, suitable(rows))
    wb.active.title = "Danh sách phù hợp"
    fill_sheet(wb.create_sheet("Tất cả CV"), rows)
    wb.save(out)


def write_csv(rows: list[dict], path):
    import csv
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:     # utf-8-sig để Excel đọc đúng tiếng Việt
        keys = [k for k, _, _ in COLUMNS]
        w = csv.DictWriter(fh, fieldnames=keys, extrasaction="ignore")
        w.writeheader()
        w.writerows({k: csv_safe(r[k]) for k in keys if k in r} for r in rows)
